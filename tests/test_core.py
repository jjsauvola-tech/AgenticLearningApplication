import io
import json
import os
import tempfile
import threading
import unittest
import urllib.request
import urllib.error
import zipfile
from pathlib import Path
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from ala.storage import Store
from ala.importers import extract, ImportProblem
from ala.server import AppServer

def docx_bytes(text='Photosynthesis converts light into chemical energy.'):
    doc=Document()
    doc.add_heading('Biology',0)
    doc.add_heading('Photosynthesis',1)
    doc.add_paragraph(text)
    table=doc.add_table(rows=1,cols=2)
    table.cell(0,0).text='Input'
    table.cell(0,1).text='Light'
    out=io.BytesIO();doc.save(out);return out.getvalue()

def pptx_bytes():
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6])
    group=slide.shapes.add_group_shape()
    nested=group.shapes.add_group_shape()
    shape=nested.shapes.add_textbox(Inches(1),Inches(1),Inches(4),Inches(2))
    shape.text='A concept inside nested groups'
    slide.notes_slide.notes_text_frame.text='Teacher-only guidance'
    out=io.BytesIO();deck.save(out);return out.getvalue()

def pdf_bytes(encrypted=False):
    writer=PdfWriter();writer.add_blank_page(width=100,height=100)
    if encrypted:writer.encrypt('password')
    out=io.BytesIO();writer.write(out);return out.getvalue()

class CoreTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.store=Store(self.tmp.name);self.vid=self.store.config['active']
    def tearDown(self):self.tmp.cleanup()
    def test_docx_and_table(self):
        result=extract('lecture.docx',docx_bytes())
        self.assertIn('chemical energy',result['text'])
        self.assertTrue(any(b['type']=='table' for p in result['pages'] for b in p['blocks']))
    def test_nested_powerpoint_groups_and_notes(self):
        result=extract('lecture.pptx',pptx_bytes())
        self.assertIn('nested groups',result['pages'][0]['text'])
        self.assertEqual('Teacher-only guidance',result['pages'][0]['notes'])
        self.assertNotIn('Teacher-only guidance',result['text'])
    def test_pdf_empty_text_is_reported(self):
        result=extract('scan.pdf',pdf_bytes())
        self.assertIn('ocr_needed',result['warnings'])
    def test_encrypted_and_corrupt_files_fail(self):
        for name,data in [('secret.pdf',pdf_bytes(True)),('bad.pptx',b'not a zip'),('file.exe',b'abcd')]:
            with self.assertRaises((ImportProblem,ValueError)):extract(name,data)
    def test_duplicate_and_new_version_preserve_notes(self):
        data=docx_bytes()
        a=self.store.add_document(self.vid,'biology.docx',data,'Biology')
        self.store.save_note(self.vid,{'document_id':a['id'],'page':1,'title':'My insight','body':'Sunlight'})
        duplicate=self.store.add_document(self.vid,'renamed.docx',data,'Biology')
        self.assertEqual(a['id'],duplicate['id'])
        self.assertTrue(duplicate['duplicate'])
        b=self.store.add_document(self.vid,'biology.docx',docx_bytes('Updated content'),'Biology')
        self.assertNotEqual(a['id'],b['id'])
        self.assertEqual(a['id'],self.store.rows(self.vid,'notes')[0]['document_id'])
    def test_independent_vaults_and_restart(self):
        self.store.add_document(self.vid,'biology.docx',docx_bytes(),'Biology')
        other=self.store.create_vault('History')
        self.assertEqual([],self.store.documents(other))
        self.store.settings({'language':'en','theme':'dark','fontSize':20,'showAssistant':False})
        restarted=Store(self.tmp.name)
        self.assertEqual('en',restarted.config['settings']['language'])
        self.assertEqual(other,restarted.config['active'])
        self.assertEqual(1,len(restarted.documents(self.vid)))
    def test_backup_restore_to_new_root(self):
        original=docx_bytes()
        a=self.store.add_document(self.vid,'biology.docx',original,'Biology')
        self.store.save_note(self.vid,{'document_id':a['id'],'page':1,'title':'Insight','body':'Remember this'})
        data=self.store.export(self.vid)
        with tempfile.TemporaryDirectory() as elsewhere:
            other=Store(elsewhere)
            restored=other.restore(data,'Restored biology')
            self.assertEqual('Remember this',other.rows(restored,'notes')[0]['body'])
            self.assertEqual(original,other.original(restored,a['id']).read_bytes())
    def test_zip_slip_restore_rejected(self):
        out=io.BytesIO()
        with zipfile.ZipFile(out,'w') as z:
            z.writestr('manifest.json','{"schema":1}')
            z.writestr('learning.sqlite','fake')
            z.writestr('../outside.txt','bad')
        with self.assertRaises(ValueError):self.store.restore(out.getvalue(),'bad')
        self.assertFalse((Path(self.tmp.name)/'outside.txt').exists())
    def test_search_and_settings_validation(self):
        a=self.store.add_document(self.vid,'biology.docx',docx_bytes(),'Biology')
        self.assertEqual(a['id'],self.store.search(self.vid,'photosynthesis')[0]['document_id'])
        self.assertEqual([],self.store.search(self.vid,"' OR 1=1 --"))
        with self.assertRaises(ValueError):self.store.settings({'ollamaUrl':'https://example.org'})
        with self.assertRaises(ValueError):self.store.settings({'language':'invalid'})

class HttpTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.server=AppServer(self.tmp.name,Path(__file__).resolve().parents[1]/'web')
        self.thread=threading.Thread(target=self.server.serve_forever,daemon=True);self.thread.start()
        self.base='http://127.0.0.1:'+str(self.server.server_port)
    def tearDown(self):
        self.server.shutdown();self.server.server_close();self.thread.join();self.tmp.cleanup()
    def request(self,path,body=None,auth=True,origin=None):
        headers={'Content-Type':'application/json'}
        if auth:headers['X-ALA-Token']=self.server.token
        if origin:headers['Origin']=origin
        req=urllib.request.Request(self.base+path,data=None if body is None else json.dumps(body).encode(),headers=headers)
        return urllib.request.urlopen(req)
    def test_unauthorized_and_cross_origin_rejected(self):
        for kwargs in [{'auth':False},{'origin':'https://malicious.example'}]:
            with self.assertRaises(urllib.error.HTTPError) as error:self.request('/api/state',**kwargs)
            self.assertEqual(403,error.exception.code)
    def test_settings_and_static_ui(self):
        with self.request('/api/settings',{'language':'en','theme':'dark'}) as r:self.assertEqual('en',json.load(r)['language'])
        with self.request('/') as r:self.assertIn(b'ALA',r.read())
    def test_model_not_configured_honest_error(self):
        with self.assertRaises(urllib.error.HTTPError) as error:self.request('/api/chat',{'text':'Explain'})
        self.assertEqual('model_not_configured',json.load(error.exception)['error'])

if __name__=='__main__':unittest.main()

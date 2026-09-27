"""Local, non-executing document import. Source bytes are retained separately."""
import io
import zipfile
from pathlib import Path

MAX_FILE = 100 * 1024 * 1024
MAX_EXPANDED = 300 * 1024 * 1024

class ImportProblem(ValueError):
    pass

def check_archive(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
            if len(infos) > 15000 or sum(i.file_size for i in infos) > MAX_EXPANDED:
                raise ImportProblem('archive_limit')
            if any(i.flag_bits & 1 for i in infos):
                raise ImportProblem('encrypted_file')
    except zipfile.BadZipFile as e:
        raise ImportProblem('invalid_file') from e

def paragraphs(text):
    return [{'type': 'text', 'text': t} for t in text.replace('\v', '\n').splitlines() if t.strip()]

def ppt_shapes(shapes):
    for shape in shapes:
        if hasattr(shape, 'shapes'):
            yield from ppt_shapes(shape.shapes)
        if shape.has_text_frame:
            yield from paragraphs(shape.text)
        if shape.has_table:
            yield {'type': 'table', 'rows': [[c.text for c in r.cells] for r in shape.table.rows]}
        if shape.shape_type == 13:
            yield {'type': 'notice', 'text': 'image_in_original'}

def block_text(block):
    if block['type'] == 'table':
        return '\n'.join(' | '.join(row) for row in block['rows'])
    return block.get('text', '') if block['type'] != 'notice' else ''

def extract(name, data):
    if not data or len(data) > MAX_FILE:
        raise ImportProblem('file_size')
    ext = Path(name).suffix.lower()
    pages = []
    warnings = []
    if ext in ('.docx', '.pptx'):
        check_archive(data)
    if ext == '.pptx':
        from pptx import Presentation
        deck = Presentation(io.BytesIO(data))
        if len(deck.slides) > 500:
            raise ImportProblem('page_limit')
        for index, slide in enumerate(deck.slides, 1):
            blocks = list(ppt_shapes(slide.shapes))
            text = '\n'.join(block_text(b) for b in blocks)
            pages.append({'number': index, 'title': next((b['text'] for b in blocks if b['type'] == 'text'), str(index)),
                          'text': text, 'blocks': blocks,
                          'notes': slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''})
        warnings.append('structured_pptx')
    elif ext == '.docx':
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc = Document(io.BytesIO(data))
        blocks = []
        title = Path(name).stem
        def flush():
            nonlocal blocks
            if blocks:
                pages.append({'number': len(pages)+1, 'title': title, 'text': '\n'.join(block_text(b) for b in blocks),
                              'blocks': blocks, 'notes': ''})
                blocks = []
        for item in doc.element.body:
            if item.tag.endswith('}p'):
                p = Paragraph(item, doc)
                if p.style and p.style.name.startswith('Heading') and p.text.strip():
                    flush()
                    title = p.text
                if p.text.strip():
                    blocks.append({'type': 'text', 'text': p.text})
                if item.xpath('.//w:drawing'):
                    blocks.append({'type': 'notice', 'text': 'image_in_original'})
            elif item.tag.endswith('}tbl'):
                table = Table(item, doc)
                blocks.append({'type': 'table', 'rows': [[c.text for c in r.cells] for r in table.rows]})
            if sum(len(block_text(b)) for b in blocks) > 10000:
                flush()
        flush()
        warnings.append('structured_docx')
    elif ext == '.pdf':
        from pypdf import PdfReader
        if not data.startswith(b'%PDF-'):
            raise ImportProblem('invalid_file')
        pdf = PdfReader(io.BytesIO(data))
        if pdf.is_encrypted:
            raise ImportProblem('encrypted_file')
        if len(pdf.pages) > 1000:
            raise ImportProblem('page_limit')
        for index, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ''
            pages.append({'number': index, 'title': str(index), 'text': text, 'blocks': paragraphs(text), 'notes': ''})
        if any(len(p['text'].strip()) < 20 for p in pages):
            warnings.append('ocr_needed')
    else:
        raise ImportProblem('unsupported_format')
    if not pages:
        raise ImportProblem('empty_document')
    if any(b['type'] == 'notice' for p in pages for b in p['blocks']):
        warnings.append('visual_content')
    return {'format': ext[1:], 'pages': pages, 'warnings': warnings,
            'text': '\n'.join(p['text'] for p in pages)}

